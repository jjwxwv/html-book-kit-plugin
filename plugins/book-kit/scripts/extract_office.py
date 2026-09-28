#!/usr/bin/env python3
"""Extract text from .pptx / .docx into markdown so agents never open the binaries.

Usage (run from the book project root; the project is the current directory,
or pass --root <dir>):
  python3 "${CLAUDE_PLUGIN_ROOT}/scripts/extract_office.py" "sources/ch2/2.1 deck.pptx" [more paths...]
Output: .book-state/extracted-office/<sha256 first 12>.md  (one per input, cached by SHA)
Relative source paths resolve against the project root.

Requires: pip install python-pptx python-docx   (only the one you need)
PDF / Markdown / text sources do NOT need this script — read them directly.
"""
import argparse
import hashlib
import json
import os
import sys

# Project resolution (plugin edition): see scan_sources.py — the project is the
# current directory (or --root / BOOK_ROOT); globals are assigned by configure().
ROOT = OUT_DIR = None


def resolve_root(explicit=None) -> str:
    return os.path.abspath(explicit or os.environ.get("BOOK_ROOT") or os.getcwd())


def configure(root: str) -> None:
    global ROOT, OUT_DIR
    ROOT = root
    OUT_DIR = os.path.join(ROOT, ".book-state", "extracted-office")


def require_book_project(root: str) -> None:
    if os.path.isfile(os.path.join(root, "book.config.json")) or os.path.isdir(os.path.join(root, "sources")):
        return
    print(json.dumps({"error": "not a book project", "root": root,
                      "hint": "run from the project root (needs book.config.json or sources/), "
                              "pass --root <dir>, or run /book-kit:init to create one"},
                     ensure_ascii=False), file=sys.stderr)
    sys.exit(2)


def sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _iter_pptx_shapes(shapes, group_type):
    """Yield leaf shapes, recursing into groups (grouped text boxes are common
    in decks and are silently lost without this)."""
    for shape in shapes:
        if shape.shape_type == group_type:
            yield from _iter_pptx_shapes(shape.shapes, group_type)
        else:
            yield shape


def pptx_to_md(path: str) -> str:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        sys.exit("python-pptx is not installed. Run: pip install python-pptx")
    prs = Presentation(path)
    lines = [f"# PPTX extraction: {os.path.basename(path)}", ""]
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"## Slide {i}")
        pictures = charts = 0
        for shape in _iter_pptx_shapes(slide.shapes, MSO_SHAPE_TYPE.GROUP):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_chart", False) and shape.has_chart:
                charts += 1
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
                    lines.append("| " + " | ".join(cells) + " |")
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(("  " * min(para.level, 4)) + "- " + text)
        if pictures or charts:
            what = " + ".join(x for x in (
                f"{pictures} picture(s)" if pictures else "",
                f"{charts} chart(s)" if charts else "") if x)
            lines.append(f"- [visual] {what} on this slide — images and chart renderings "
                         "are NOT extracted from PPTX; infer meaning from surrounding text "
                         "only, and flag a gap if the visual itself carries the teaching "
                         "content (ask the user for a PDF export of image-heavy decks)")
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        if notes:
            lines.append(f"> notes: {notes}")
        lines.append("")
    return "\n".join(lines)


def docx_to_md(path: str) -> str:
    try:
        import docx
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        sys.exit("python-docx is not installed. Run: pip install python-docx")
    d = docx.Document(path)
    lines = [f"# DOCX extraction: {os.path.basename(path)}", ""]
    t_i = 0
    # Walk the body in document order so tables keep their position in the
    # flow (appending them at the end loses the heading/context they belong to).
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, d)
            text = (para.text or "").strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            if style.startswith("heading"):
                try:
                    level = int(style.split()[-1])
                except ValueError:
                    level = 2
                lines.append("#" * min(level + 1, 6) + " " + text)
            else:
                lines.append(text)
            lines.append("")
        elif child.tag == qn("w:tbl"):
            t_i += 1
            table = Table(child, d)
            lines.append(f"### Table {t_i}")
            for row in table.rows:
                cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    return "\n".join(lines)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="+", help=".pptx / .docx source files (relative to the project root)")
    ap.add_argument("--root", help="book project root (default: current directory or $BOOK_ROOT)")
    args = ap.parse_args()
    root = resolve_root(args.root)
    require_book_project(root)
    configure(root)
    os.makedirs(OUT_DIR, exist_ok=True)
    for arg in args.paths:
        path = arg if os.path.isabs(arg) else os.path.join(ROOT, arg)
        if not os.path.isfile(path):
            print(f"SKIP (not found): {arg}")
            continue
        ext = os.path.splitext(path)[1].lower()
        digest = sha256(path)
        out = os.path.join(OUT_DIR, digest[:12] + ".md")
        if os.path.isfile(out):
            print(f"CACHED: {arg} -> {os.path.relpath(out, ROOT)}")
            continue
        if ext == ".pptx":
            md = pptx_to_md(path)
        elif ext == ".docx":
            md = docx_to_md(path)
        else:
            print(f"SKIP (only .pptx/.docx need extraction): {arg}")
            continue
        header = f"<!-- source: {os.path.relpath(path, ROOT)} | sha256: {digest} -->\n\n"
        with open(out, "w", encoding="utf-8") as f:
            f.write(header + md)
        print(f"OK: {arg} -> {os.path.relpath(out, ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
